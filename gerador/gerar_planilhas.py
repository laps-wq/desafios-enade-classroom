# -*- coding: utf-8 -*-
"""Gera as planilhas a partir de calendario.py, questoes_exemplo.json e, se existir, questoes.json.

Uso (na raiz do repositório):  python3 gerador/gerar_planilhas.py

Sempre gera (vão para o repositório público):
  kit-professores/Planilha_Modelo_Desafios.xlsx  planilha-modelo com o desafio de exemplo EX01
  kit-professores/Modelo_Envio_Questoes.xlsx     modelo para professores enviarem questões

Só gera quando gerador/questoes.json existe (banco com gabarito, fora do repositório):
  planilhas/Desafios_ENADE_TC2.xlsx              planilha de controle da TC2/TC3
  questoes/banco_questoes_TC2.md                 banco de questões legível
  e acrescenta ao modelo de envio a aba "Questões atuais".
O .gitignore impede que esses arquivos sejam enviados ao GitHub.
"""
import json
import os
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from calendario import AR, DESAFIOS

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
def caminho(*p): return os.path.join(RAIZ, *p)

# ---------------------------------------------------------------------
# Banco de questões
# ---------------------------------------------------------------------
EXEMPLOS = json.load(open(caminho("gerador", "questoes_exemplo.json"), encoding="utf-8"))
ARQ_BANCO = caminho("gerador", "questoes.json")
TEM_BANCO = os.path.exists(ARQ_BANCO)
QUESTOES = json.load(open(ARQ_BANCO, encoding="utf-8")) if TEM_BANCO else []
QUESTOES.sort(key=lambda q: (q["desafio"], q["n"]))

ids = [d[0] for d in DESAFIOS]
por_desafio = defaultdict(list)
for q in QUESTOES:
    por_desafio[q["desafio"]].append(q)
for q in QUESTOES + EXEMPLOS:
    assert 2 <= len(q["alternativas"]) <= 5 and q["correta"] in "ABCDE"[:len(q["alternativas"])], (q["desafio"], q["n"])
if TEM_BANCO:
    assert set(por_desafio) == set(ids), "todo desafio do calendário precisa de questões"
    print("Banco:", len(QUESTOES), "questões · gabarito:", dict(sorted(Counter(q["correta"] for q in QUESTOES).items())))
else:
    print("gerador/questoes.json não encontrado: gerando só os arquivos do kit (sem o banco de questões).")

# ---------------------------------------------------------------------
# Estilo (cores da CESAR School)
# ---------------------------------------------------------------------
ESCURO, LARANJA, CREME = "232323", "FF6002", "FEF7EE"
CAB_FONT, CAB_FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=ESCURO)
BORDA = Border(bottom=Side(style="thin", color="DDDDDD"))
TOPO = Alignment(vertical="top", wrap_text=True)
LISTRA = PatternFill("solid", fgColor=CREME)
CORES_STATUS = {"PRONTO": "FFF2CC", "AGENDADO": "DDEBF7", "PUBLICADO": "E2EFDA", "ENCERRADO": "EDEDED", "ERRO": "F8CBAD"}

def cabecalho(ws, colunas, larguras):
    ws.append(colunas)
    for i, c in enumerate(ws[1], start=1):
        c.font, c.fill = CAB_FONT, CAB_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = larguras[i - 1]
    ws.row_dimensions[1].height = 30

def texto(ws, linhas, largura=120):
    ws.column_dimensions["A"].width = largura
    ws.sheet_view.showGridLines = False
    for t, tipo in linhas:
        ws.append([t])
        c = ws.cell(row=ws.max_row, column=1)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if tipo == "titulo":
            c.font = Font(bold=True, size=15, color=LARANJA); ws.row_dimensions[ws.max_row].height = 26
        elif tipo == "secao":
            c.font = Font(bold=True, size=11, color="FFFFFF"); c.fill = CAB_FILL; ws.row_dimensions[ws.max_row].height = 20
        elif tipo == "alerta":
            c.font = Font(bold=True, color="9C3D00")

def lista(ws, intervalo, opcoes):
    v = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
    ws.add_data_validation(v); v.add(intervalo)

def estilizar(ws, negrito=(0,), correta=None, listrar_por=None):
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.alignment = TOPO; c.border = BORDA
        for i in negrito:
            row[i].font = Font(bold=True)
        if correta is not None:
            row[correta].font = Font(bold=True, color=LARANJA)
            row[correta].alignment = Alignment(horizontal="center", vertical="top")
        if listrar_por is not None and row[listrar_por].value and int(str(row[listrar_por].value)[-2:]) % 2 == 0:
            for c in row:
                c.fill = LISTRA

def aba_desafios(wb, linhas):
    ws = wb.create_sheet("Desafios")
    cols = ["ID", "Título", "Temas", "Aulas cobertas", "Publicar em", "Prazo", "Pontos", "Status",
            "Mensagens do robô", "Link do formulário", "Link de edição", "Link no Classroom",
            "ID do formulário", "IDs no Classroom"]
    cabecalho(ws, cols, [7, 42, 34, 52, 17, 17, 8, 13, 52, 30, 30, 30, 22, 26])
    for l in linhas:
        ws.append(l + [10, "RASCUNHO", "", "", "", "", "", ""])
    for r in range(2, 60):
        for col in (5, 6):
            ws.cell(row=r, column=col).number_format = "dd/mm/yyyy hh:mm"
        for col in (13, 14):
            ws.cell(row=r, column=col).number_format = "@"
        if r <= len(linhas) + 1:
            for col in range(1, len(cols) + 1):
                c = ws.cell(row=r, column=col)
                c.alignment = Alignment(vertical="top", wrap_text=col in (2, 3, 4, 9)); c.border = BORDA
            ws.cell(row=r, column=1).font = Font(bold=True)
    lista(ws, "H2:H200", ["RASCUNHO", "PRONTO", "AGENDADO", "PUBLICADO", "ENCERRADO", "ERRO"])
    for st, cor in CORES_STATUS.items():
        ws.conditional_formatting.add("H2:H200", FormulaRule(formula=[f'$H2="{st}"'], fill=PatternFill("solid", fgColor=cor)))
    ws.freeze_panes = "C2"

def aba_questoes(wb, questoes):
    ws = wb.create_sheet("Questoes")
    cols = ["Desafio", "Nº", "Área", "Enunciado", "A", "B", "C", "D", "E", "Correta", "Comentário", "Pontos", "Imagem"]
    cabecalho(ws, cols, [9, 5, 22, 72, 34, 34, 34, 34, 34, 9, 62, 8, 20])
    for desafio, n, q in questoes:
        alts = q["alternativas"] + [""] * (5 - len(q["alternativas"]))
        ws.append([desafio, n, q["area"], q["enunciado"], *alts, q["correta"], q["comentario"], 2, ""])
    estilizar(ws, correta=9, listrar_por=0 if questoes and questoes[0][0].startswith("D") else None)
    lista(ws, "J2:J500", list("ABCDE"))
    ws.freeze_panes = "D2"

def aba_config(wb, texto_extra):
    ws = wb.create_sheet("Config")
    cabecalho(ws, ["Chave", "Valor", "O que é"], [28, 40, 100])
    for k, v, d in [
        ("ID_TURMAS", "", "ID da turma no Classroom (menu > Listar minhas turmas). Para mais de uma turma, separe por vírgula."),
        ("DOMINIO_ALUNOS", "cesar.school", "Domínio do e-mail dos alunos. Só contas desse domínio conseguem abrir o formulário."),
        ("RESPONDENTES", "DOMINIO", "DOMINIO (qualquer conta do domínio acima com o link) · ALUNOS (só os alunos das turmas) · QUALQUER (qualquer pessoa com o link)."),
        ("TOPICO_CLASSROOM", "Desafios da Semana", "Tópico do mural de Atividades onde os desafios e gabaritos ficam agrupados (é criado se não existir)."),
        ("NOTAS_NO_CLASSROOM", "RASCUNHO", "RASCUNHO (você revisa e devolve) · ATRIBUIDA (lança direto como nota atribuída) · NAO (só registra na aba Notas)."),
        ("PUBLICAR_GABARITO", "SIM", "Publica o gabarito comentado no Classroom logo após o prazo de cada desafio."),
        ("FECHAR_NO_PRAZO", "SIM", "Fecha o formulário no prazo (até 15 min depois). Com NÃO, ele continua aceitando respostas atrasadas."),
        ("UMA_RESPOSTA_POR_ALUNO", "SIM", "Cada aluno só pode enviar o formulário uma vez."),
        ("EMBARALHAR_QUESTOES", "NÃO", "Embaralha a ordem das questões para cada aluno."),
        ("AVISAR_ERROS_POR_EMAIL", "SIM", "Quando o piloto automático encontra um erro, manda um e-mail para você."),
        ("PASTA_DRIVE_ID", "", "Opcional: link ou ID de uma pasta do Drive onde guardar os formulários criados."),
        ("FORM_MODELO_ID", "", "Opcional: link de um Forms-modelo cujas configurações serão copiadas (ex.: esconder o gabarito até o prazo)."),
        ("TEXTO_EXTRA_NA_ATIVIDADE", texto_extra, "Opcional: texto acrescentado ao final da descrição de toda atividade no Classroom."),
    ]:
        ws.append([k, v, d])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[0].font = Font(bold=True); row[1].number_format = "@"; row[1].fill = LISTRA
        for c in row:
            c.alignment = TOPO; c.border = BORDA
    lista(ws, "B4", ["DOMINIO", "ALUNOS", "QUALQUER"]); lista(ws, "B6", ["RASCUNHO", "ATRIBUIDA", "NAO"])
    for r in range(7, 12):
        lista(ws, f"B{r}", ["SIM", "NÃO"])
    ws.freeze_panes = "A2"

def abas_notas_log(wb):
    cabecalho(wb.create_sheet("Notas"), ["Desafio", "E-mail", "Pontos no formulário", "Máximo", "Nota", "Enviado em", "Atualizado em"], [9, 34, 18, 9, 8, 18, 18])
    cabecalho(wb.create_sheet("Log"), ["Quando", "Ação", "Desafio", "Detalhe"], [18, 12, 9, 100])

INSTALACAO = [
    ("INSTALAÇÃO (uma vez só, uns 10 minutos)", "secao"),
    ("1. No Google Drive da conta @cesar.school: Novo > Upload deste arquivo .xlsx. Abra e use Arquivo > Salvar como Planilhas Google.", ""),
    ("2. Na planilha Google: Extensões > Apps Script. Apague o conteúdo de Código.gs e cole o conteúdo do arquivo Codigo.gs.", ""),
    ("3. No Apps Script: Configurações do projeto (engrenagem) > marque 'Mostrar arquivo de manifesto appsscript.json no editor'.", ""),
    ("   Volte ao Editor, abra appsscript.json, apague tudo e cole o conteúdo do arquivo appsscript.json. Salve (Ctrl+S).", ""),
    ("4. Recarregue a planilha. Vai aparecer o menu 'Desafios ENADE'. Rode '1) Verificar configuração' e autorize o acesso.", ""),
    ("5. Rode '2) Listar minhas turmas do Classroom'. Copie o ID da turma e cole em Config > ID_TURMAS (várias turmas: separe por vírgula).", ""),
    ("6. Confira DOMINIO_ALUNOS na aba Config (o domínio do e-mail dos alunos, ex.: cesar.school).", ""),
]
STATUS_E_DICAS = [
    ("", ""),
    ("O QUE ACONTECE COM CADA DESAFIO", "secao"),
    ("RASCUNHO → em revisão; o script ignora.", ""),
    ("PRONTO → o script cria o Google Forms (modo teste, com correção automática) e agenda a atividade no Classroom.", ""),
    ("AGENDADO → a atividade aparece para os alunos sozinha no horário de 'Publicar em' (o próprio Classroom publica).", ""),
    ("PUBLICADO → aberta para respostas até o 'Prazo'.", ""),
    ("ENCERRADO → no prazo, o formulário fecha, as notas vão para o Classroom como RASCUNHO e o gabarito comentado é publicado.", ""),
    ("ERRO → a coluna 'Mensagens do robô' explica o problema. Corrija e volte o Status para PRONTO. Você também recebe um e-mail.", ""),
    ("", ""),
    ("DICAS", "secao"),
    ("• Para trocar datas, edite 'Publicar em' e 'Prazo' (formato dd/mm/aaaa hh:mm) antes de marcar PRONTO.", ""),
    ("• Errou algo num desafio já AGENDADO (que ainda não saiu)? Selecione a linha e use 'Refazer o desafio selecionado'.", ""),
    ("• Desafio já publicado com erro em uma questão: edite direto no formulário pelo 'Link de edição'. As respostas são mantidas.", ""),
    ("• As notas entram como rascunho: no Classroom, abra a atividade e clique em 'Devolver' para os alunos verem.", ""),
    ("• Imagens: cole na coluna Imagem da questão o link de um arquivo do seu Drive; ela aparece antes do enunciado.", ""),
    ("• Para os alunos não verem o gabarito logo ao enviar: crie um Forms vazio em modo teste, em Configurações > Testes", ""),
    ("  desmarque 'Respostas corretas' (ou escolha liberar a nota depois), e cole o link dele em Config > FORM_MODELO_ID.", ""),
]

# ---------------------------------------------------------------------
# 1. Planilha da TC2/TC3
# ---------------------------------------------------------------------
if TEM_BANCO:
  wb = Workbook()
  texto(wb.active, [
    ("Desafios Semanais ENADE · Tópicos Contemporâneos 2 e 3 · 2026.2", "titulo"),
    ("Dois desafios por semana, publicados no Classroom às 12:30 de segunda e quarta. Cada um fica aberto até o próximo sair.", ""),
    ("", ""),
] + INSTALACAO + [
    ("", ""),
    ("ANTES DE LIGAR", "secao"),
    ("7. Na aba Desafios, clique na linha D01 e rode '3) Pré-visualizar'. Abra o link 'Ver como aluno' e responda como teste.", ""),
    ("   Se puder, peça para um aluno ou monitor abrir o link e confirmar que tem acesso.", ""),
    ("8. Revise as questões na aba Questoes. Quando um desafio estiver aprovado, mude o Status dele de RASCUNHO para PRONTO.", ""),
    ("9. Menu > 'Ligar piloto automático'. Pronto: a cada 15 minutos o script agenda os PRONTOS e encerra os vencidos.", ""),
  ] + STATUS_E_DICAS)
  wb.active.title = "Como usar"
  aba_desafios(wb, [[i, t, temas, aulas, pub, prazo] for (i, t, temas, aulas, pub, prazo) in DESAFIOS])
  aba_questoes(wb, [(q["desafio"], q["n"], q) for q in QUESTOES])
  aba_config(wb, "")
  abas_notas_log(wb)
  os.makedirs(caminho("planilhas"), exist_ok=True)
  wb.save(caminho("planilhas", "Desafios_ENADE_TC2.xlsx"))

# ---------------------------------------------------------------------
# 2. Planilha-modelo genérica (kit)
# ---------------------------------------------------------------------
wb = Workbook()
texto(wb.active, [
    ("Desafios no Google Classroom · planilha-modelo", "titulo"),
    ("Você escreve as questões e as datas. O script cria um Google Forms em modo teste, agenda a atividade no Classroom e,", ""),
    ("no prazo, fecha o formulário, lança as notas como rascunho e publica o gabarito comentado.", ""),
    ("", ""),
] + INSTALACAO + [
    ("", ""),
    ("MONTANDO OS SEUS DESAFIOS", "secao"),
    ("7. Na aba Desafios, uma linha por desafio: ID (ex.: D01), título, temas, 'Publicar em' e 'Prazo' (dd/mm/aaaa hh:mm) e pontos.", ""),
    ("8. Na aba Questoes, uma linha por questão, com o mesmo ID na coluna Desafio. De 2 a 5 alternativas, a letra correta e um comentário.", ""),
    ("   A linha EX01 é um exemplo com três formatos de questão (direta, I/II/III e asserção-razão). Apague ou edite quando quiser.", ""),
    ("9. Clique na linha de um desafio e rode '3) Pré-visualizar'. Abra 'Ver como aluno' e peça a um aluno para testar o link.", ""),
    ("10. Desafio revisado? Mude o Status de RASCUNHO para PRONTO.", ""),
    ("11. Menu > 'Ligar piloto automático'. A cada 15 minutos o script agenda os PRONTOS e encerra os vencidos.", ""),
] + STATUS_E_DICAS)
wb.active.title = "Como usar"
aba_desafios(wb, [["EX01", "Desafio 1 · Exemplo de formatos de questão", "Exemplo: SO, Redes e Teoria da Computação",
                   "Troque pelo tema e pela aula da sua disciplina", datetime(2026, 10, 19, 12, 30), datetime(2026, 10, 21, 12, 30)]])
aba_questoes(wb, [("EX01", q["n"], q) for q in EXEMPLOS])  # direta, I/II/III e asserção-razão
aba_config(wb, "Questões no estilo ENADE. Conta para a nota de participação.")
abas_notas_log(wb)
wb.save(caminho("kit-professores", "Planilha_Modelo_Desafios.xlsx"))

# ---------------------------------------------------------------------
# 3. Modelo de envio de questões (kit)
# ---------------------------------------------------------------------
if TEM_BANCO:
    intro = [
        ("Cada aula de conteúdo da TC2/TC3 tem um desafio de 5 questões no Classroom. As questões já estão escritas;", ""),
        ("este arquivo serve para você revisar as da sua aula e mandar correções ou questões novas.", ""),
        ("Atenção: este arquivo contém o gabarito. Não compartilhe com alunos.", "alerta"),
        ("", ""),
        ("PASSO A PASSO", "secao"),
        ("1. Na aba Calendário, encontre o ID do desafio da sua aula (ex.: D06 = Sistemas Operacionais, sai em 21/10).", ""),
        ("2. Na aba Questões atuais, use o filtro da coluna Desafio para ver as 5 questões desse ID.", ""),
    ]
else:
    intro = [
        ("Cada aula de conteúdo da TC2/TC3 tem um desafio de 5 questões no Classroom. Este arquivo serve para você", ""),
        ("mandar questões novas ou corrigidas para a sua aula.", ""),
        ("As questões atuais não estão nesta versão do arquivo, porque têm gabarito. Para revisá-las, peça à Laura a versão com o banco.", "alerta"),
        ("", ""),
        ("PASSO A PASSO", "secao"),
        ("1. Na aba Calendário, encontre o ID do desafio da sua aula (ex.: D06 = Sistemas Operacionais, sai em 21/10).", ""),
        ("2. Se for revisar as questões atuais, use a versão com o banco que a Laura enviar.", ""),
    ]
wb = Workbook()
texto(wb.active, [
    ("Desafios ENADE · TC2 e TC3 · revisão e envio de questões", "titulo"),
] + intro + [
    ("3. Na aba Minhas questões, escreva uma linha por questão nova ou corrigida:", ""),
    ("   • para corrigir ou trocar uma questão atual, coloque o número dela na coluna 'Substitui a nº';", ""),
    ("   • para acrescentar, deixe essa coluna vazia.", ""),
    ("4. Devolva este arquivo para a Laura, de preferência até uma semana antes da sua aula.", ""),
    ("", ""),
    ("COMO ESCREVER CADA QUESTÃO", "secao"),
    ("• Enunciado com contexto (uma situação, um trecho de código, uma tabela), como no ENADE.", ""),
    ("• 5 alternativas (A a E), só uma correta. Evite 'todas as anteriores' e 'nenhuma das anteriores'.", ""),
    ("• Coluna Correta: só a letra.", ""),
    ("• Comentário: 2 a 4 frases explicando por que a correta está certa e onde estão as pegadinhas. O aluno vê isso no gabarito.", ""),
    ("• Imagem (opcional): link de um arquivo no Google Drive, compartilhado com a Laura. Ela aparece antes do enunciado.", ""),
    ("", ""),
    ("ALTERNATIVAS PADRÃO DO ENADE", "secao"),
    ("Questões do tipo 'É correto o que se afirma em': use combinações como I, apenas · II, apenas · I e III, apenas · II e III, apenas · I, II e III.", ""),
    ("Questões de asserção-razão (I. ... PORQUE II. ...) usam sempre estas cinco alternativas:", ""),
] + [(f"{l}) {t}", "") for l, t in zip("ABCDE", AR)], largura=130)
wb.active.title = "Como preencher"

ws = wb.create_sheet("Calendário")
cabecalho(ws, ["ID", "Sai em", "Prazo", "Tema", "Aulas cobertas"], [7, 17, 17, 44, 90])
for (i, t, temas, aulas, pub, prazo) in DESAFIOS:
    ws.append([i, pub, prazo, t.split("· ", 1)[1], aulas])
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    row[1].number_format = row[2].number_format = "ddd dd/mm hh:mm"
estilizar(ws, listrar_por=0)
ws.freeze_panes = "B2"

if TEM_BANCO:
    ws = wb.create_sheet("Questões atuais")
    cabecalho(ws, ["Desafio", "Nº", "Área", "Enunciado", "A", "B", "C", "D", "E", "Correta", "Comentário"], [9, 5, 22, 72, 34, 34, 34, 34, 34, 9, 62])
    for q in QUESTOES:
        alts = q["alternativas"] + [""] * (5 - len(q["alternativas"]))
        ws.append([q["desafio"], q["n"], q["area"], q["enunciado"], *alts, q["correta"], q["comentario"]])
    estilizar(ws, correta=9, listrar_por=0)
    ws.auto_filter.ref = f"A1:K{ws.max_row}"
    ws.freeze_panes = "D2"

ws = wb.create_sheet("Minhas questões")
cols = ["Desafio", "Substitui a nº", "Área", "Enunciado", "A", "B", "C", "D", "E", "Correta", "Comentário", "Imagem", "Professor(a)"]
cabecalho(ws, cols, [9, 10, 22, 72, 34, 34, 34, 34, 34, 9, 62, 20, 20])
for r in range(2, 32):
    for col in range(1, len(cols) + 1):
        c = ws.cell(row=r, column=col); c.alignment = TOPO; c.border = BORDA
    ws.row_dimensions[r].height = 60
lista(ws, "A2:A200", ids); lista(ws, "B2:B200", ["1", "2", "3", "4", "5"]); lista(ws, "J2:J200", list("ABCDE"))
ws.freeze_panes = "D2"
wb.active = 0
wb.save(caminho("kit-professores", "Modelo_Envio_Questoes.xlsx"))

# ---------------------------------------------------------------------
# 4. Banco legível em Markdown
# ---------------------------------------------------------------------
if TEM_BANCO:
    md = ["# Desafios Semanais ENADE · TC2/TC3 — banco de questões", "",
          "Gerado por `gerador/gerar_planilhas.py` a partir de `gerador/questoes.json` (16 desafios × 5 questões, 2 pontos cada).", ""]
    for (i, t, temas, aulas, pub, prazo) in DESAFIOS:
        md += [f"## {i} · {t.split('· ', 1)[1]}", f"Publica {pub:%d/%m %H:%M} · prazo {prazo:%d/%m %H:%M} · {aulas}", ""]
        for q in por_desafio[i]:
            md += [f"**Questão {q['n']} ({q['area']})**", "", q["enunciado"], ""]
            md += [f"- {l}) {a}" for l, a in zip("ABCDE", q["alternativas"])]
            md += ["", f"Gabarito: **{q['correta']}** · {q['comentario']}", ""]
    os.makedirs(caminho("questoes"), exist_ok=True)
    open(caminho("questoes", "banco_questoes_TC2.md"), "w", encoding="utf-8").write("\n".join(md))
    print("OK: planilhas/, kit-professores/ e questoes/ atualizados (com o banco)")
else:
    print("OK: kit-professores/ atualizado (sem o banco)")
